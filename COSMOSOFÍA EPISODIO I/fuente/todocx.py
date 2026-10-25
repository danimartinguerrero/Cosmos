import sys
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
NAVY=RGBColor(0x1b,0x2f,0x4e); BLUE=RGBColor(0x3a,0x64,0xa0); GREY=RGBColor(0x6b,0x74,0x80)
soup=BeautifulSoup(open(sys.argv[1],encoding='utf-8'),'html.parser')
d=Document()
st=d.styles['Normal']; st.font.name='Palatino Linotype'; st.font.size=Pt(11.5)
for s in d.sections: s.left_margin=s.right_margin=Cm(2.4)
def para(text='',size=None,color=None,bold=False,italic=False,align=None,after=6,before=0,indent=None):
    p=d.add_paragraph(); r=p.add_run(text); r.bold=bold; r.italic=italic
    if size: r.font.size=Pt(size)
    if color: r.font.color.rgb=color
    if align is not None: p.alignment=align
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(before)
    if indent: p.paragraph_format.left_indent=Cm(indent)
    return p
para('RADIOACCIÓN CÓSMICA · ASOCIACIÓN SUPERNOVA',9,GREY,align=WD_ALIGN_PARAGRAPH.CENTER)
para('COSMOSOFÍA',30,NAVY,bold=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=2)
para('Episodio I · Un mar de cuatrocientos mil millones de soles',14,NAVY,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER)
para('Guion confeccionado por Daniel Martín Guerrero',11,GREY,align=WD_ALIGN_PARAGRAPH.CENTER,after=14)
for sec in soup.select('#script .sec'):
    n=sec.select_one('h2 .n').get_text(); title=sec.select_one('h2').get_text()[len(n):]
    para(f'{n}. {title.upper()}',13,NAVY,bold=True,before=14,after=0)
    para(sec.select_one('.tag').get_text(),10.5,GREY,italic=True,after=8)
    for el in sec.find_all(['div','figure'],recursive=False):
        c=el.get('class') or []
        if 'l' in c:
            who=el.select_one('.who').get_text(); txt=' '.join(el.p.get_text().split())
            p=d.add_paragraph(); r=p.add_run(who+': '); r.bold=True; r.font.color.rgb=BLUE
            p.add_run(txt); p.paragraph_format.space_after=Pt(6)
        elif 'cue' in c:
            para('['+el.get_text().strip().upper()+']',9.5,GREY,after=6)
        elif 'ref' in c:
            para(el.select_one('.h').get_text(' ',strip=True),9.5,BLUE,bold=True,indent=1.2,after=1)
            para(' '.join(el.blockquote.get_text().split()),11,NAVY,italic=True,indent=1.2,after=1)
            para('— '+el.select_one('.src').get_text(' ',strip=True),9,GREY,indent=1.2,after=8)
para('FUENTES Y CITAS',13,NAVY,bold=True,before=18)
for i,li in enumerate(soup.select('.sources ol li'),1):
    para(f'{i}. '+' '.join(li.get_text().split()),9.5,after=3)
para(' '.join(soup.select_one('.sources .note').get_text().split()),9.5,GREY,italic=True,before=8)
d.save(sys.argv[2])
